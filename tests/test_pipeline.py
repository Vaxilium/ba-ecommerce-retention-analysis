"""Pipeline tests: generator reproducibility, SQL translation, statistics, artifacts.

Fast unit tests always run; artifact tests skip if outputs/ has not been built yet
(CI runs `make all` before pytest, so they execute there).
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
PY = sys.executable
AN = ROOT / "outputs" / "analysis"


def _hashes(d: Path) -> dict:
    return {p.name: hashlib.md5(p.read_bytes()).hexdigest() for p in sorted(d.glob("*.csv"))}


class TestGenerator:
    def test_same_seed_is_reproducible(self, tmp_path):
        outs = []
        for run in ("a", "b"):
            out = tmp_path / f"out_{run}"
            subprocess.run([PY, str(ROOT / "src" / "generate_data.py"), "--customers", "300",
                            "--products", "60", "--seed", "7", "--out", str(out),
                            "--samples-dir", str(tmp_path / f"smp_{run}")],
                           cwd=tmp_path, check=True, capture_output=True, timeout=300)
            outs.append(out)
        assert _hashes(outs[0]) == _hashes(outs[1])

    def test_different_seeds_differ(self, tmp_path):
        for seed, name in ((7, "s7"), (8, "s8")):
            out = tmp_path / name
            subprocess.run([PY, str(ROOT / "src" / "generate_data.py"), "--customers", "300",
                            "--products", "60", "--seed", str(seed), "--out", str(out),
                            "--samples-dir", str(tmp_path / f"smp_{name}")],
                           cwd=tmp_path, check=True, capture_output=True, timeout=300)
        assert _hashes(tmp_path / "s7") != _hashes(tmp_path / "s8")

    def test_committed_samples_exist(self):
        for name in ("orders", "customers", "order_items", "products",
                     "support_tickets", "marketing_touches"):
            assert (ROOT / "data" / "sample" / f"{name}_sample.csv").exists()


class TestSqlTranslation:
    def test_date_literal(self):
        from build_warehouse import translate_for_sqlite
        assert translate_for_sqlite("WHERE d <= DATE '2025-12-31'") == "WHERE d <= '2025-12-31'"

    def test_date_diff_days(self):
        from build_warehouse import translate_for_sqlite
        out = translate_for_sqlite("SELECT date_diff('day', a, b) FROM t")
        assert "julianday(b) - julianday(a)" in out

    def test_timestamp_cast(self):
        from build_warehouse import translate_for_sqlite
        assert "TIMESTAMP)" not in translate_for_sqlite("CAST(x AS TIMESTAMP)")


class TestStatistics:
    def test_z_critical_values(self):
        from ab_test_winback import z_two_sided
        assert z_two_sided(0.05) == pytest.approx(1.959964, abs=1e-4)
        assert z_two_sided(0.10) == pytest.approx(1.644854, abs=1e-4)

    def test_p_value_round_trip(self):
        from ab_test_winback import normal_p_value, z_two_sided
        z = z_two_sided(0.05)
        assert normal_p_value(z) == pytest.approx(0.05, abs=1e-5)
        assert normal_p_value(0.0) == pytest.approx(1.0)
        assert normal_p_value(-z) == pytest.approx(normal_p_value(z))


@pytest.fixture(scope="module")
def summary():
    return json.loads((AN / "summary.json").read_text())


@pytest.mark.skipif(not (AN / "summary.json").exists(), reason="run make all first")
class TestArtifacts:
    def test_kpi_decline_consistent(self, summary):
        k = summary["kpi"]
        assert k["repeat_rate_h1_2024"] > k["repeat_rate_q3_2025"]
        assert k["decline_pp"] == pytest.approx(
            (k["repeat_rate_q3_2025"] - k["repeat_rate_h1_2024"]) * 100, abs=0.1)

    def test_decomposition_sums(self, summary):
        d = summary["decomposition"]["joint_channel_x_discount"]
        assert d["mix_pp"] + d["rate_pp"] == pytest.approx(d["total_pp"], abs=0.05)

    def test_channel_rates_bounded(self, summary):
        for row in summary["channel"]:
            assert 0 < row["repeat_rate_90d"] < 1
            assert row["customers"] > 0

    def test_winback_decision_matches_statistics(self, summary):
        ab = json.loads((AN / "ab_test.json").read_text())
        o = ab["observed"]
        if o["p_value"] < ab["design"]["alpha"] and o["abs_lift"] > 0:
            assert ab["decision"] == "ROLL OUT"
        assert o["treatment_rate"] != o["control_rate"]

    def test_rfm_revenue_shares_sum(self, summary):
        prof = json.loads((AN / "rfm_segments.json").read_text())["profiles"]
        assert sum(p["pct_revenue"] for p in prof) == pytest.approx(100, abs=1.0)
        assert sum(p["customers"] for p in prof) == json.loads(
            (AN / "rfm_segments.json").read_text())["total_customers"]

    def test_business_case_workbook_has_live_formulas(self):
        from openpyxl import load_workbook
        wb = load_workbook(ROOT / "outputs" / "business_case.xlsx")
        ws = wb["Scenarios"]
        roi_cell = ws.cell(row=14, column=2).value
        assert isinstance(roi_cell, str) and roi_cell.startswith("=")
        assert isinstance(wb["Assumptions"].cell(row=3, column=2).value, (int, float))

    def test_churn_model_metrics_sane(self):
        m = json.loads((AN / "churn_model.json").read_text())
        assert 0.5 < m["auc_test"] < 1.0
        assert m["n_test"] > 0

    def test_input_ordering_is_deterministic(self):
        """Queries feeding the split are ORDER BY customer_id; run model twice is cheap only with
        warehouse present."""
        if not (ROOT / "outputs" / "warehouse.duckdb").exists():
            pytest.skip("warehouse not built")
        import duckdb
        con = duckdb.connect(str(ROOT / "outputs" / "warehouse.duckdb"), read_only=True)
        a = con.execute("SELECT customer_id FROM dim_customer ORDER BY customer_id LIMIT 5").fetchall()
        b = con.execute("SELECT customer_id FROM dim_customer ORDER BY customer_id LIMIT 5").fetchall()
        con.close()
        assert a == b
